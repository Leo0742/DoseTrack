import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from PIL import Image
from pillow_heif import register_heif_opener

from app.models import Treatment
from app.services.intakes import create_regimen, ensure_intakes_for_date, resolve_intake


def _jpeg() -> bytes:
    payload = io.BytesIO()
    Image.new("RGB", (32, 32), "white").save(payload, format="JPEG")
    return payload.getvalue()


def _heic() -> bytes:
    register_heif_opener()
    payload = io.BytesIO()
    Image.new("RGB", (32, 32), "white").save(payload, format="HEIF", quality=90)
    return payload.getvalue()


def _png() -> bytes:
    payload = io.BytesIO()
    Image.new("RGBA", (32, 32), (255, 255, 255, 128)).save(payload, format="PNG")
    return payload.getvalue()


def test_heic_photo_is_converted_to_jpeg(owner_client):
    upload = owner_client.post(
        "/api/photos",
        data={"photo_date": "2026-09-22", "title": "iPhone HEIC"},
        files={"file": ("iphone.heic", _heic(), "image/heic")},
    )
    assert upload.status_code == 200

    content = owner_client.get(f"/api/photos/{upload.json()['id']}/content")
    assert content.status_code == 200
    assert content.headers["content-type"].startswith("image/jpeg")
    assert content.content.startswith(b"\xff\xd8\xff")


def test_png_photo_stays_png(owner_client):
    upload = owner_client.post(
        "/api/photos",
        data={"photo_date": "2026-09-22", "title": "iPhone PNG"},
        files={"file": ("iphone.png", _png(), "image/png")},
    )
    assert upload.status_code == 200

    content = owner_client.get(f"/api/photos/{upload.json()['id']}/content")
    assert content.status_code == 200
    assert content.headers["content-type"].startswith("image/png")
    assert content.content.startswith(b"\x89PNG\r\n\x1a\n")


def test_photo_folder_lifecycle(owner_client):
    folder_response = owner_client.post("/api/photo-folders", data={"name": "Лицо"})
    assert folder_response.status_code == 200
    folder_id = folder_response.json()["id"]

    upload = owner_client.post(
        "/api/photos",
        data={"photo_date": "2026-09-21", "folder_id": folder_id, "title": "День 1"},
        files={"file": ("progress.jpg", _jpeg(), "image/jpeg")},
    )
    assert upload.status_code == 200
    photo_id = upload.json()["id"]
    assert upload.json()["folder_id"] == folder_id

    deleted_folder = owner_client.delete(f"/api/photo-folders/{folder_id}")
    assert deleted_folder.status_code == 200
    photos = owner_client.get("/api/photos").json()
    assert len(photos) == 1
    assert photos[0]["id"] == photo_id
    assert photos[0]["folder_id"] is None

    assert owner_client.delete(f"/api/photos/{photo_id}").status_code == 200
    assert owner_client.get("/api/photos").json() == []


def test_nested_photo_folders_and_move(owner_client):
    parent = owner_client.post("/api/photo-folders", data={"name": "2026"}).json()
    child = owner_client.post(
        "/api/photo-folders", data={"name": "Сентябрь", "parent_id": parent["id"]}
    ).json()
    folders = owner_client.get("/api/photo-folders").json()
    assert next(row for row in folders if row["id"] == child["id"])["parent_id"] == parent["id"]

    uploaded = owner_client.post(
        "/api/photos",
        data={"photo_date": "2026-09-21", "title": "Фото"},
        files={"file": ("progress.jpg", _jpeg(), "image/jpeg")},
    ).json()
    moved = owner_client.put(
        f"/api/photos/{uploaded['id']}/folder", data={"folder_id": child["id"]}
    )
    assert moved.status_code == 200
    assert moved.json()["folder_id"] == child["id"]

    assert owner_client.delete(f"/api/photo-folders/{parent['id']}").status_code == 200
    folders_after = owner_client.get("/api/photo-folders").json()
    assert next(row for row in folders_after if row["id"] == child["id"])["parent_id"] is None


def test_medication_excel_export_ru_and_en(db, owner, owner_client):
    treatment = Treatment(
        patient_id=owner.id,
        name="Course",
        medication_name="Isotretinoin",
        start_date=date(2026, 9, 21),
        status="ACTIVE",
        target_mg=Decimal("120"),
        timezone="Europe/Moscow",
    )
    db.add(treatment)
    db.commit()
    db.refresh(treatment)
    create_regimen(
        db,
        treatment,
        owner,
        treatment.start_date,
        [{"slot_key": "morning", "label": "Morning", "planned_time": "08:00", "dose_mg": "20"}],
    )
    intake = ensure_intakes_for_date(db, treatment, treatment.start_date)[0]
    resolve_intake(db, intake, owner, "TAKEN")
    next_day = date(2026, 9, 22)
    skipped = ensure_intakes_for_date(db, treatment, next_day)[0]
    resolve_intake(db, skipped, owner, "SKIPPED", reason="Forgot")

    ru = owner_client.get("/api/export/medication.xlsx?lang=ru")
    assert ru.status_code == 200
    ru_book = load_workbook(io.BytesIO(ru.content), read_only=True)
    assert ru_book.sheetnames == ["Итоги", "История приёма", "Вес", "Схемы приёма"]
    assert ru_book["Итоги"]["A1"].value == "Пациент"
    assert ru_book["История приёма"]["A1"].value == "Дата"
    summary = {row[0].value: row[1].value for row in ru_book["Итоги"].iter_rows(min_col=1, max_col=2)}
    assert summary["Принято за период, мг"] == 20
    assert summary["Пропущено приёмов"] == 1
    assert summary["Осталось к цели на конец периода, мг"] == 100

    ranged = owner_client.get(
        "/api/export/medication.xlsx?lang=ru&start=2026-09-22&end=2026-09-22"
    )
    assert ranged.status_code == 200
    ranged_book = load_workbook(io.BytesIO(ranged.content), read_only=True)
    ranged_summary = {
        row[0].value: row[1].value
        for row in ranged_book["Итоги"].iter_rows(min_col=1, max_col=2)
    }
    assert ranged_summary["Принято до периода, мг"] == 20
    assert ranged_summary["Принято за период, мг"] == 0
    assert ranged_summary["Пропущено приёмов"] == 1
    history_dates = [row[0].value for row in ranged_book["История приёма"].iter_rows(min_row=2)]
    assert [value.date() if hasattr(value, "date") else value for value in history_dates] == [next_day]

    en = owner_client.get("/api/export/medication.xlsx?lang=en")
    assert en.status_code == 200
    en_book = load_workbook(io.BytesIO(en.content), read_only=True)
    assert en_book.sheetnames == ["Summary", "Medication history", "Weight history", "Regimen history"]
    assert en_book["Summary"]["A1"].value == "Patient"
    assert en_book["Medication history"]["A1"].value == "Date"
