import csv
import io
from datetime import UTC, date, datetime, time, timedelta
from typing import cast
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from app.api.schemas import DiaryCommentInput, DiaryInput, WeightInput
from app.core.deps import current_user, require_patient_write, resolve_patient_id
from app.db import get_db
from app.models import (
    DiaryComment,
    DiaryEntry,
    Document,
    DocumentCategory,
    DocumentFolder,
    DoseRegimen,
    PhotoFolder,
    ProgressPhoto,
    ScheduledIntake,
    Treatment,
    User,
    WeightRecord,
)
from app.services.audit import audit
from app.services.forecast import current_daily_planned, estimate_completion
from app.services.intakes import ensure_intakes_range
from app.services.notifications import queue_doctor_event
from app.services.storage import (
    delete_storage,
    private_path,
    read_validated_upload,
    store_private,
    store_progress_photo,
    thumbnail_path,
)

router = APIRouter(tags=["records"])
DEFAULT_DOCUMENT_CATEGORIES = [
    "Blood tests",
    "Doctor visits",
    "Prescriptions",
    "Medical reports",
    "Other",
]


@router.get("/weights")
def weights(user: User = Depends(current_user), db: Session = Depends(get_db)):
    patient_id = resolve_patient_id(db, user)
    rows = db.scalars(
        select(WeightRecord)
        .where(WeightRecord.patient_id == patient_id)
        .order_by(WeightRecord.recorded_date.desc())
    )
    return [
        {
            "id": row.id,
            "date": row.recorded_date,
            "weight_kg": float(row.weight_kg),
            "note": row.note,
        }
        for row in rows
    ]


@router.post("/weights")
def save_weight(
    payload: WeightInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = db.scalar(
        select(WeightRecord).where(
            WeightRecord.patient_id == patient_id,
            WeightRecord.recorded_date == payload.recorded_date,
        )
    )
    before = None
    if row:
        before = {"weight_kg": str(row.weight_kg), "note": row.note}
        row.weight_kg = payload.weight_kg
        row.note = payload.note
    else:
        row = WeightRecord(
            patient_id=patient_id,
            recorded_date=payload.recorded_date,
            weight_kg=payload.weight_kg,
            note=payload.note,
            created_by=user.id,
        )
        db.add(row)
    db.flush()
    audit(
        db,
        user,
        "weight_changed",
        "weight_record",
        row.id,
        {"before": before, "after": {"weight_kg": str(payload.weight_kg)}},
    )
    db.commit()
    return {
        "id": row.id,
        "date": row.recorded_date,
        "weight_kg": float(row.weight_kg),
        "note": row.note,
    }


@router.get("/diary")
def diary(
    category: str | None = None,
    q: str | None = None,
    day: date | None = Query(default=None, alias="date"),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    stmt = select(DiaryEntry).where(DiaryEntry.patient_id == patient_id)
    if day:
        patient = db.get(User, patient_id)
        timezone = ZoneInfo(patient.timezone if patient else user.timezone)
        day_start = datetime.combine(day, time.min, tzinfo=timezone).astimezone(UTC)
        next_day_start = datetime.combine(day + timedelta(days=1), time.min, tzinfo=timezone).astimezone(UTC)
        stmt = stmt.where(
            DiaryEntry.occurred_at >= day_start,
            DiaryEntry.occurred_at < next_day_start,
        )
    if category:
        stmt = stmt.where(DiaryEntry.category == category)
    if q:
        stmt = stmt.where((DiaryEntry.body.ilike(f"%{q}%")) | (DiaryEntry.title.ilike(f"%{q}%")))
    rows = db.scalars(stmt.order_by(DiaryEntry.occurred_at.desc()).limit(500))
    return [diary_dict(row, db, user) for row in rows]


@router.post("/diary")
def add_diary(
    payload: DiaryInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = DiaryEntry(
        patient_id=patient_id,
        occurred_at=payload.occurred_at,
        title=payload.title,
        body=payload.body,
        category=payload.category,
        severity=payload.severity,
        created_by=user.id,
    )
    db.add(row)
    db.flush()
    audit(db, user, "diary_created", "diary_entry", row.id)
    db.commit()
    return diary_dict(row, db, user)


@router.put("/diary/{entry_id}")
def edit_diary(
    entry_id: str,
    payload: DiaryInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = db.get(DiaryEntry, entry_id)
    if not row or row.patient_id != patient_id:
        raise HTTPException(status_code=404, detail="Diary entry not found")
    before = jsonable_encoder(diary_dict(row))
    row.occurred_at = payload.occurred_at
    row.title = payload.title
    row.body = payload.body
    row.category = payload.category
    row.severity = payload.severity
    audit(db, user, "diary_edited", "diary_entry", row.id, {"before": before})
    db.commit()
    return diary_dict(row, db, user)


@router.delete("/diary/{entry_id}")
def delete_diary(
    entry_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = db.get(DiaryEntry, entry_id)
    if not row or row.patient_id != patient_id:
        raise HTTPException(status_code=404, detail="Diary entry not found")
    audit(db, user, "diary_deleted", "diary_entry", row.id)
    db.delete(row)
    db.commit()
    return {"ok": True}


def diary_dict(row: DiaryEntry, db: Session | None = None, user: User | None = None) -> dict:
    payload = {
        "id": row.id,
        "occurred_at": row.occurred_at,
        "title": row.title,
        "body": row.body,
        "category": row.category,
        "severity": row.severity,
        "created_at": row.created_at,
        "updated_at": row.updated_at,
    }
    if db is not None and user is not None:
        comments = list(
            db.scalars(
                select(DiaryComment)
                .where(DiaryComment.entry_id == row.id)
                .order_by(DiaryComment.created_at)
            )
        )
        payload["comments"] = [comment_dict(db, comment, user) for comment in comments]
    return payload


def comment_dict(db: Session, row: DiaryComment, user: User) -> dict:
    author = db.get(User, row.author_id)
    name = ""
    if author:
        name = " ".join(part for part in (author.first_name, author.last_name) if part).strip()
        if not name:
            name = author.username
    return {
        "id": row.id,
        "entry_id": row.entry_id,
        "body": row.body,
        "author_id": row.author_id,
        "author_name": name,
        "author_role": author.role if author else "",
        "is_mine": row.author_id == user.id,
        "created_at": row.created_at,
        "updated_at": row.updated_at,
    }


@router.post("/diary/{entry_id}/comments")
def add_diary_comment(
    entry_id: str,
    payload: DiaryCommentInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    entry = db.get(DiaryEntry, entry_id)
    if not entry or entry.patient_id != patient_id:
        raise HTTPException(status_code=404, detail="Diary entry not found")
    row = DiaryComment(
        entry_id=entry.id,
        patient_id=patient_id,
        author_id=user.id,
        body=payload.body.strip(),
    )
    db.add(row)
    db.flush()
    audit(db, user, "diary_comment_created", "diary_comment", row.id)
    db.commit()
    db.refresh(row)
    return comment_dict(db, row, user)


@router.put("/diary/comments/{comment_id}")
def edit_diary_comment(
    comment_id: str,
    payload: DiaryCommentInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    row = db.get(DiaryComment, comment_id)
    if not row or row.patient_id != patient_id:
        raise HTTPException(status_code=404, detail="Comment not found")
    if row.author_id != user.id:
        raise HTTPException(status_code=403, detail="Only the author can edit this comment")
    row.body = payload.body.strip()
    audit(db, user, "diary_comment_edited", "diary_comment", row.id)
    db.commit()
    db.refresh(row)
    return comment_dict(db, row, user)


@router.delete("/diary/comments/{comment_id}")
def delete_diary_comment(
    comment_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    row = db.get(DiaryComment, comment_id)
    if not row or row.patient_id != patient_id:
        raise HTTPException(status_code=404, detail="Comment not found")
    if row.author_id != user.id:
        raise HTTPException(status_code=403, detail="Only the author can delete this comment")
    audit(db, user, "diary_comment_deleted", "diary_comment", row.id)
    db.delete(row)
    db.commit()
    return {"ok": True}


@router.get("/document-categories")
def document_categories(user: User = Depends(current_user), db: Session = Depends(get_db)):
    patient_id = resolve_patient_id(db, user)
    rows = list(
        db.scalars(
            select(DocumentCategory)
            .where(DocumentCategory.patient_id == patient_id)
            .order_by(DocumentCategory.name)
        )
    )
    if not rows and user.id == patient_id:
        rows = [DocumentCategory(patient_id=patient_id, name=name) for name in DEFAULT_DOCUMENT_CATEGORIES]
        db.add_all(rows)
        db.commit()
    return [{"id": row.id, "name": row.name} for row in rows]


@router.post("/document-categories")
def create_document_category(
    name: str = Form(...),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    clean = name.strip()[:120]
    if not clean:
        raise HTTPException(status_code=422, detail="Category name is required")
    existing = db.scalar(
        select(DocumentCategory).where(
            DocumentCategory.patient_id == patient_id,
            DocumentCategory.name == clean,
        )
    )
    if existing:
        return {"id": existing.id, "name": existing.name}
    row = DocumentCategory(patient_id=patient_id, name=clean)
    db.add(row)
    db.commit()
    db.refresh(row)
    return {"id": row.id, "name": row.name}


def _validate_document_parent(db: Session, patient_id: str, parent_id: str | None) -> None:
    if not parent_id:
        return
    parent = db.get(DocumentFolder, parent_id)
    if not parent or parent.patient_id != patient_id:
        raise HTTPException(status_code=422, detail="Invalid document folder parent")


@router.get("/document-folders")
def document_folders(user: User = Depends(current_user), db: Session = Depends(get_db)):
    patient_id = resolve_patient_id(db, user)
    folders = list(
        db.scalars(
            select(DocumentFolder)
            .where(DocumentFolder.patient_id == patient_id)
            .order_by(DocumentFolder.name)
        )
    )
    docs = list(
        db.scalars(
            select(Document).where(Document.patient_id == patient_id, Document.deleted_at.is_(None))
        )
    )
    item_counts: dict[str, int] = {}
    child_counts: dict[str, int] = {}
    for doc in docs:
        if doc.folder_id:
            item_counts[doc.folder_id] = item_counts.get(doc.folder_id, 0) + 1
    for folder in folders:
        if folder.parent_id:
            child_counts[folder.parent_id] = child_counts.get(folder.parent_id, 0) + 1
    return [
        {
            "id": folder.id,
            "name": folder.name,
            "parent_id": folder.parent_id,
            "item_count": item_counts.get(folder.id, 0),
            "child_count": child_counts.get(folder.id, 0),
        }
        for folder in folders
    ]


@router.post("/document-folders")
def create_document_folder(
    name: str = Form(...),
    parent_id: str | None = Form(default=None),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    clean = name.strip()[:120]
    if not clean:
        raise HTTPException(status_code=422, detail="Folder name is required")
    _validate_document_parent(db, patient_id, parent_id)
    duplicate = db.scalar(
        select(DocumentFolder).where(
            DocumentFolder.patient_id == patient_id,
            DocumentFolder.parent_id == parent_id,
            func.lower(DocumentFolder.name) == clean.lower(),
        )
    )
    if duplicate:
        return {"id": duplicate.id, "name": duplicate.name, "parent_id": duplicate.parent_id}
    row = DocumentFolder(patient_id=patient_id, parent_id=parent_id, name=clean)
    db.add(row)
    db.flush()
    audit(db, user, "document_folder_created", "document_folder", row.id, {"name": clean})
    db.commit()
    db.refresh(row)
    return {"id": row.id, "name": row.name, "parent_id": row.parent_id}


@router.put("/document-folders/{folder_id}")
def rename_document_folder(
    folder_id: str,
    name: str = Form(...),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = db.get(DocumentFolder, folder_id)
    if not row or row.patient_id != patient_id:
        raise HTTPException(status_code=404, detail="Document folder not found")
    clean = name.strip()[:120]
    if not clean:
        raise HTTPException(status_code=422, detail="Folder name is required")
    row.name = clean
    audit(db, user, "document_folder_renamed", "document_folder", row.id, {"name": clean})
    db.commit()
    return {"id": row.id, "name": row.name, "parent_id": row.parent_id}


@router.delete("/document-folders/{folder_id}")
def delete_document_folder(
    folder_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = db.get(DocumentFolder, folder_id)
    if not row or row.patient_id != patient_id:
        raise HTTPException(status_code=404, detail="Document folder not found")
    db.execute(
        update(Document)
        .where(Document.patient_id == patient_id, Document.folder_id == row.id)
        .values(folder_id=row.parent_id)
    )
    db.execute(
        update(DocumentFolder)
        .where(DocumentFolder.patient_id == patient_id, DocumentFolder.parent_id == row.id)
        .values(parent_id=row.parent_id)
    )
    audit(db, user, "document_folder_deleted", "document_folder", row.id, {"name": row.name})
    db.delete(row)
    db.commit()
    return {"ok": True}


@router.get("/documents")
def documents(
    q: str | None = None,
    category_id: str | None = None,
    start: date | None = None,
    end: date | None = None,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    stmt = select(Document).where(Document.patient_id == patient_id, Document.deleted_at.is_(None))
    if q:
        stmt = stmt.where(Document.title.ilike(f"%{q}%"))
    if category_id:
        stmt = stmt.where(Document.category_id == category_id)
    if start:
        stmt = stmt.where(Document.document_date >= start)
    if end:
        stmt = stmt.where(Document.document_date <= end)
    rows = db.scalars(stmt.order_by(Document.document_date.desc(), Document.created_at.desc()))
    return [document_dict(row) for row in rows]


@router.post("/documents")
async def upload_document(
    file: UploadFile = File(...),
    title: str = Form(...),
    category_id: str | None = Form(default=None),
    folder_id: str | None = Form(default=None),
    document_date: date | None = Form(default=None),
    note: str | None = Form(default=None),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    data, mime = await read_validated_upload(file)
    if category_id:
        category = db.get(DocumentCategory, category_id)
        if not category or category.patient_id != patient_id:
            raise HTTPException(status_code=422, detail="Invalid document category")
    _validate_document_parent(db, patient_id, folder_id)
    suffix = ".pdf" if mime == "application/pdf" else ".bin"
    storage_id, size, checksum = store_private(data, suffix)
    row = Document(
        patient_id=patient_id,
        category_id=category_id,
        folder_id=folder_id,
        original_filename=(file.filename or "document")[:255],
        storage_id=storage_id,
        title=title[:180],
        document_date=document_date,
        note=note,
        mime_type=mime,
        size_bytes=size,
        checksum_sha256=checksum,
        uploader_id=user.id,
    )
    db.add(row)
    db.flush()
    audit(db, user, "document_uploaded", "document", row.id, {"mime_type": mime, "size": size})
    queue_doctor_event(
        db,
        patient_id,
        "document_uploaded",
        "document",
        row.id,
        checksum,
        {"title": row.title, "document_id": row.id, "document_date": str(document_date or "")},
    )
    db.commit()
    return document_dict(row)


@router.get("/documents/{document_id}/content")
def document_content(
    document_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    row = db.get(Document, document_id)
    if not row or row.patient_id != patient_id or row.deleted_at:
        raise HTTPException(status_code=404, detail="Document not found")
    audit(db, user, "document_viewed", "document", row.id)
    db.commit()
    return FileResponse(
        private_path(row.storage_id),
        media_type=row.mime_type,
        filename=row.original_filename,
        content_disposition_type="inline",
    )


@router.delete("/documents/{document_id}")
def delete_document(
    document_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = db.get(Document, document_id)
    if not row or row.patient_id != patient_id or row.deleted_at:
        raise HTTPException(status_code=404, detail="Document not found")
    row.deleted_at = datetime.now(UTC)
    delete_storage(row.storage_id)
    audit(db, user, "document_deleted", "document", row.id)
    db.commit()
    return {"ok": True}


def document_dict(row: Document) -> dict:
    return {
        "id": row.id,
        "title": row.title,
        "original_filename": row.original_filename,
        "category_id": row.category_id,
        "folder_id": row.folder_id,
        "document_date": row.document_date,
        "mime_type": row.mime_type,
        "size_bytes": row.size_bytes,
        "checksum_sha256": row.checksum_sha256,
        "note": row.note,
        "created_at": row.created_at,
    }


@router.put("/documents/{document_id}/folder")
def move_document(
    document_id: str,
    folder_id: str | None = Form(default=None),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = db.get(Document, document_id)
    if not row or row.patient_id != patient_id or row.deleted_at:
        raise HTTPException(status_code=404, detail="Document not found")
    _validate_document_parent(db, patient_id, folder_id)
    row.folder_id = folder_id
    audit(db, user, "document_moved", "document", row.id, {"folder_id": folder_id})
    db.commit()
    return document_dict(row)


@router.get("/photo-folders")
def photo_folders(user: User = Depends(current_user), db: Session = Depends(get_db)):
    patient_id = resolve_patient_id(db, user)
    rows = list(
        db.scalars(
            select(PhotoFolder)
            .where(PhotoFolder.patient_id == patient_id)
            .order_by(PhotoFolder.name)
        )
    )
    photos = list(
        db.scalars(
            select(ProgressPhoto).where(
                ProgressPhoto.patient_id == patient_id,
                ProgressPhoto.deleted_at.is_(None),
            )
        )
    )
    item_counts: dict[str, int] = {}
    child_counts: dict[str, int] = {}
    for photo in photos:
        if photo.folder_id:
            item_counts[photo.folder_id] = item_counts.get(photo.folder_id, 0) + 1
    for row in rows:
        if row.parent_id:
            child_counts[row.parent_id] = child_counts.get(row.parent_id, 0) + 1
    return [
        {
            "id": row.id,
            "name": row.name,
            "parent_id": row.parent_id,
            "item_count": item_counts.get(row.id, 0),
            "child_count": child_counts.get(row.id, 0),
        }
        for row in rows
    ]


@router.post("/photo-folders")
def create_photo_folder(
    name: str = Form(...),
    parent_id: str | None = Form(default=None),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    clean = name.strip()[:120]
    if not clean:
        raise HTTPException(status_code=422, detail="Folder name is required")
    if parent_id:
        parent = db.get(PhotoFolder, parent_id)
        if not parent or parent.patient_id != patient_id:
            raise HTTPException(status_code=422, detail="Invalid photo folder parent")
    existing = db.scalar(
        select(PhotoFolder).where(
            PhotoFolder.patient_id == patient_id,
            PhotoFolder.name == clean,
        )
    )
    if existing:
        return {"id": existing.id, "name": existing.name, "parent_id": existing.parent_id}
    row = PhotoFolder(patient_id=patient_id, parent_id=parent_id, name=clean)
    db.add(row)
    db.flush()
    audit(db, user, "photo_folder_created", "photo_folder", row.id, {"name": clean})
    db.commit()
    db.refresh(row)
    return {"id": row.id, "name": row.name, "parent_id": row.parent_id}


@router.put("/photo-folders/{folder_id}")
def rename_photo_folder(
    folder_id: str,
    name: str = Form(...),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    folder = db.get(PhotoFolder, folder_id)
    if not folder or folder.patient_id != patient_id:
        raise HTTPException(status_code=404, detail="Photo folder not found")
    clean = name.strip()[:120]
    if not clean:
        raise HTTPException(status_code=422, detail="Folder name is required")
    folder.name = clean
    audit(db, user, "photo_folder_renamed", "photo_folder", folder.id, {"name": clean})
    db.commit()
    return {"id": folder.id, "name": folder.name, "parent_id": folder.parent_id}


@router.delete("/photo-folders/{folder_id}")
def delete_photo_folder(
    folder_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    folder = db.get(PhotoFolder, folder_id)
    if not folder or folder.patient_id != patient_id:
        raise HTTPException(status_code=404, detail="Photo folder not found")
    db.execute(
        update(ProgressPhoto)
        .where(ProgressPhoto.patient_id == patient_id, ProgressPhoto.folder_id == folder.id)
        .values(folder_id=folder.parent_id)
    )
    db.execute(
        update(PhotoFolder)
        .where(PhotoFolder.patient_id == patient_id, PhotoFolder.parent_id == folder.id)
        .values(parent_id=folder.parent_id)
    )
    audit(db, user, "photo_folder_deleted", "photo_folder", folder.id, {"name": folder.name})
    db.delete(folder)
    db.commit()
    return {"ok": True}


@router.get("/photos")
def photos(user: User = Depends(current_user), db: Session = Depends(get_db)):
    patient_id = resolve_patient_id(db, user)
    rows = db.scalars(
        select(ProgressPhoto)
        .where(ProgressPhoto.patient_id == patient_id, ProgressPhoto.deleted_at.is_(None))
        .order_by(ProgressPhoto.photo_date.desc(), ProgressPhoto.created_at.desc())
    )
    return [photo_dict(row) for row in rows]


@router.post("/photos")
async def upload_photo(
    file: UploadFile = File(...),
    photo_date: date = Form(...),
    folder_id: str | None = Form(default=None),
    title: str | None = Form(default=None),
    note: str | None = Form(default=None),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    if folder_id:
        folder = db.get(PhotoFolder, folder_id)
        if not folder or folder.patient_id != patient_id:
            raise HTTPException(status_code=422, detail="Invalid photo folder")
    data, mime = await read_validated_upload(file, image_only=True)
    storage_id, thumb_id, size, checksum, stored_mime = store_progress_photo(data, mime)
    row = ProgressPhoto(
        patient_id=patient_id,
        folder_id=folder_id,
        photo_date=photo_date,
        angle="other",
        title=title[:180] if title else None,
        note=note,
        storage_id=storage_id,
        thumbnail_id=thumb_id,
        mime_type=stored_mime,
        size_bytes=size,
        checksum_sha256=checksum,
        uploader_id=user.id,
    )
    db.add(row)
    db.flush()
    audit(db, user, "photo_uploaded", "progress_photo", row.id)
    queue_doctor_event(
        db,
        patient_id,
        "photo_uploaded",
        "progress_photo",
        row.id,
        checksum,
        {"title": row.title or "Progress photo", "photo_id": row.id, "photo_date": str(row.photo_date)},
    )
    db.commit()
    return photo_dict(row)


@router.get("/photos/{photo_id}/content")
def photo_content(photo_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    patient_id = resolve_patient_id(db, user)
    row = db.get(ProgressPhoto, photo_id)
    if not row or row.patient_id != patient_id or row.deleted_at:
        raise HTTPException(status_code=404, detail="Photo not found")
    return FileResponse(private_path(row.storage_id), media_type=row.mime_type)


@router.get("/photos/{photo_id}/thumbnail")
def photo_thumbnail(
    photo_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    patient_id = resolve_patient_id(db, user)
    row = db.get(ProgressPhoto, photo_id)
    if not row or row.patient_id != patient_id or row.deleted_at:
        raise HTTPException(status_code=404, detail="Photo not found")
    return FileResponse(thumbnail_path(row.thumbnail_id), media_type="image/jpeg")


@router.delete("/photos/{photo_id}")
def delete_photo(photo_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = db.get(ProgressPhoto, photo_id)
    if not row or row.patient_id != patient_id or row.deleted_at:
        raise HTTPException(status_code=404, detail="Photo not found")
    row.deleted_at = datetime.now(UTC)
    delete_storage(row.storage_id, row.thumbnail_id)
    audit(db, user, "photo_deleted", "progress_photo", row.id)
    db.commit()
    return {"ok": True}


def photo_dict(row: ProgressPhoto) -> dict:
    return {
        "id": row.id,
        "photo_date": row.photo_date,
        "month": row.photo_date.strftime("%Y-%m"),
        "folder_id": row.folder_id,
        "angle": row.angle,
        "title": row.title,
        "note": row.note,
        "size_bytes": row.size_bytes,
        "created_at": row.created_at,
    }


@router.put("/photos/{photo_id}/folder")
def move_photo(
    photo_id: str,
    folder_id: str | None = Form(default=None),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    require_patient_write(db, user, patient_id)
    row = db.get(ProgressPhoto, photo_id)
    if not row or row.patient_id != patient_id or row.deleted_at:
        raise HTTPException(status_code=404, detail="Photo not found")
    if folder_id:
        folder = db.get(PhotoFolder, folder_id)
        if not folder or folder.patient_id != patient_id:
            raise HTTPException(status_code=422, detail="Invalid photo folder")
    row.folder_id = folder_id
    audit(db, user, "photo_moved", "progress_photo", row.id, {"folder_id": folder_id})
    db.commit()
    return photo_dict(row)


def _safe_xlsx(value):
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


@router.get("/export/medication.xlsx")
def export_medication_xlsx(
    lang: str = Query(default="ru", pattern="^(ru|en)$"),
    start: date | None = None,
    end: date | None = None,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    if user.id != patient_id:
        raise HTTPException(status_code=403, detail="Only the owner can export data")

    treatment = db.scalar(
        select(Treatment)
        .where(Treatment.patient_id == patient_id)
        .order_by(Treatment.start_date.desc())
    )
    if not treatment:
        raise HTTPException(status_code=404, detail="No treatment to export")
    if start and end and start > end:
        raise HTTPException(status_code=422, detail="Start date must not be after end date")

    ru = lang == "ru"
    today = datetime.now(ZoneInfo(treatment.timezone)).date()
    period_start = max(start or treatment.start_date, treatment.start_date)
    period_end = min(end or today, today)
    if period_end < period_start:
        raise HTTPException(status_code=422, detail="Export period does not overlap treatment")

    ensure_intakes_range(db, treatment, treatment.start_date, period_end)
    all_rows = list(
        db.scalars(
            select(ScheduledIntake)
            .where(ScheduledIntake.treatment_id == treatment.id)
            .order_by(ScheduledIntake.scheduled_date, ScheduledIntake.planned_time)
        )
    )
    rows = [row for row in all_rows if period_start <= row.scheduled_date <= period_end]
    before_rows = [row for row in all_rows if row.scheduled_date < period_start]

    workbook = Workbook()
    summary = cast(Worksheet, workbook.active)
    summary.title = "Итоги" if ru else "Summary"
    history = workbook.create_sheet("История приёма" if ru else "Medication history")
    weights_sheet = workbook.create_sheet("Вес" if ru else "Weight history")
    regimens_sheet = workbook.create_sheet("Схемы приёма" if ru else "Regimen history")

    taken_rows = [row for row in rows if row.status == "TAKEN"]
    skipped_rows = [row for row in rows if row.status == "SKIPPED"]
    pending_rows = [row for row in rows if row.status == "PENDING"]
    cumulative_before = sum(
        (float(row.actual_dose_mg or 0) for row in before_rows if row.status == "TAKEN"),
        0.0,
    )
    period_taken_mg = sum((float(row.actual_dose_mg or 0) for row in taken_rows), 0.0)
    cumulative_end = cumulative_before + period_taken_mg
    planned_mg = sum((float(row.planned_dose_mg) for row in rows), 0.0)
    resolved = len(taken_rows) + len(skipped_rows)
    adherence = round(len(taken_rows) / resolved * 100, 1) if resolved else None
    remaining_at_end = max(float(treatment.target_mg) - cumulative_end, 0.0)
    forecast = estimate_completion(db, treatment, today)
    patient_name = " ".join(
        part for part in (user.first_name, user.last_name) if part
    ).strip() or user.username
    current_daily = float(current_daily_planned(db, treatment, today))

    labels = (
        [
            ("Пациент", patient_name),
            ("Препарат", treatment.medication_name if treatment else ""),
            ("Дата начала", treatment.start_date if treatment else ""),
            ("Период выписки", f"{period_start.isoformat()} — {period_end.isoformat()}"),
            ("Целевая суммарная доза, мг", float(treatment.target_mg) if treatment else ""),
            ("Принято до периода, мг", cumulative_before),
            ("Принято за период, мг", period_taken_mg),
            ("Накоплено к концу периода, мг", cumulative_end),
            ("Осталось к цели на конец периода, мг", remaining_at_end),
            ("Запланировано за период, мг", planned_mg),
            ("Принято приёмов", len(taken_rows)),
            ("Пропущено приёмов", len(skipped_rows)),
            ("Ожидается приёмов", len(pending_rows)),
            ("Соблюдение среди завершённых приёмов, %", adherence if adherence is not None else ""),
            ("Текущая схема, мг/день", current_daily),
            ("Ориентировочно осталось дней", forecast["estimated_days"] if forecast["estimated_days"] is not None else ""),
            ("Ориентировочная дата завершения", forecast["estimated_date"] or ""),
        ]
        if ru
        else [
            ("Patient", patient_name),
            ("Medication", treatment.medication_name if treatment else ""),
            ("Start date", treatment.start_date if treatment else ""),
            ("Statement period", f"{period_start.isoformat()} — {period_end.isoformat()}"),
            ("Cumulative target, mg", float(treatment.target_mg) if treatment else ""),
            ("Taken before period, mg", cumulative_before),
            ("Taken during period, mg", period_taken_mg),
            ("Cumulative at period end, mg", cumulative_end),
            ("Remaining to target at period end, mg", remaining_at_end),
            ("Planned during period, mg", planned_mg),
            ("Taken doses", len(taken_rows)),
            ("Skipped doses", len(skipped_rows)),
            ("Pending doses", len(pending_rows)),
            ("Adherence among resolved doses, %", adherence if adherence is not None else ""),
            ("Current regimen, mg/day", current_daily),
            ("Estimated days remaining", forecast["estimated_days"] if forecast["estimated_days"] is not None else ""),
            ("Estimated completion date", forecast["estimated_date"] or ""),
        ]
    )
    for label, value in labels:
        summary.append([label, _safe_xlsx(value)])

    headers = (
        [
            "Дата",
            "Приём",
            "Плановое время",
            "Плановая доза, мг",
            "Статус",
            "Фактическая доза, мг",
            "Фактическое время",
            "Причина пропуска",
            "Накоплено, мг",
        ]
        if ru
        else [
            "Date",
            "Dose",
            "Planned time",
            "Planned dose, mg",
            "Status",
            "Actual dose, mg",
            "Taken at",
            "Skipped reason",
            "Cumulative, mg",
        ]
    )
    history.append(headers)
    running = cumulative_before
    status_ru = {"TAKEN": "Принято", "SKIPPED": "Пропущено", "PENDING": "Ожидается"}
    for row in rows:
        if row.status == "TAKEN":
            running += float(row.actual_dose_mg or 0)
        status_value = status_ru.get(row.status, row.status) if ru else row.status.title()
        history.append(
            [
                row.scheduled_date,
                _safe_xlsx(row.label),
                row.planned_time.strftime("%H:%M"),
                float(row.planned_dose_mg),
                status_value,
                float(row.actual_dose_mg) if row.actual_dose_mg is not None else "",
                row.taken_at.replace(tzinfo=None) if row.taken_at else "",
                _safe_xlsx(row.skipped_reason or ""),
                running,
            ]
        )

    weights_sheet.append(
        ["Дата", "Вес, кг", "Заметка"] if ru else ["Date", "Weight, kg", "Note"]
    )
    weight_rows = db.scalars(
        select(WeightRecord)
        .where(
            WeightRecord.patient_id == patient_id,
            WeightRecord.recorded_date >= period_start,
            WeightRecord.recorded_date <= period_end,
        )
        .order_by(WeightRecord.recorded_date)
    )
    for weight_row in weight_rows:
        weights_sheet.append(
            [
                weight_row.recorded_date,
                float(weight_row.weight_kg),
                _safe_xlsx(weight_row.note or ""),
            ]
        )

    regimens_sheet.append(
        ["Действует с", "Действует до", "Сумма, мг/день", "Приёмы", "Заметка"]
        if ru
        else ["Effective from", "Effective to", "Total, mg/day", "Slots", "Note"]
    )
    regimen_rows = db.scalars(
        select(DoseRegimen)
        .where(DoseRegimen.treatment_id == treatment.id)
        .order_by(DoseRegimen.effective_from)
    )
    for regimen in regimen_rows:
        if regimen.effective_from > period_end:
            continue
        if regimen.effective_to is not None and regimen.effective_to < period_start:
            continue
        daily = sum((float(slot.dose_mg) for slot in regimen.slots), 0.0)
        slot_text = "; ".join(
            f"{slot.label} {slot.planned_time.strftime('%H:%M')} — {float(slot.dose_mg):g} mg"
            for slot in regimen.slots
        )
        regimens_sheet.append(
            [
                regimen.effective_from,
                regimen.effective_to or "",
                daily,
                _safe_xlsx(slot_text),
                _safe_xlsx(regimen.notes or ""),
            ]
        )

    header_fill = PatternFill("solid", fgColor="E7F0EC")
    for sheet in (summary, history, weights_sheet, regimens_sheet):
        sheet.freeze_panes = "A2" if sheet is not summary else None
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for column_index, column in enumerate(sheet.columns, start=1):
            width = min(max((len(str(cell.value or "")) for cell in column), default=8) + 2, 34)
            sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)
        for sheet_row in sheet.iter_rows():
            for cell in sheet_row:
                cell.alignment = Alignment(vertical="top")

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="dosetrack-medication-history.xlsx"'},
    )


@router.get("/export/{kind}")
def export_csv(
    kind: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    patient_id = resolve_patient_id(db, user)
    if user.id != patient_id:
        raise HTTPException(status_code=403, detail="Only the owner can export data")
    output = io.StringIO()
    writer = csv.writer(output)
    if kind == "medication":
        writer.writerow(
            [
                "date",
                "slot",
                "planned_time",
                "planned_dose_mg",
                "status",
                "actual_dose_mg",
                "taken_at",
                "skipped_reason",
            ]
        )
        rows = db.scalars(
            select(ScheduledIntake)
            .join(Treatment, Treatment.id == ScheduledIntake.treatment_id)
            .where(Treatment.patient_id == patient_id)
            .order_by(ScheduledIntake.scheduled_date, ScheduledIntake.planned_time)
        )
        for intake_row in rows:
            writer.writerow(
                [
                    intake_row.scheduled_date,
                    intake_row.label,
                    intake_row.planned_time,
                    intake_row.planned_dose_mg,
                    intake_row.status,
                    intake_row.actual_dose_mg or "",
                    intake_row.taken_at or "",
                    intake_row.skipped_reason or "",
                ]
            )
    elif kind == "weights":
        writer.writerow(["date", "weight_kg", "note"])
        for weight_row in db.scalars(
            select(WeightRecord)
            .where(WeightRecord.patient_id == patient_id)
            .order_by(WeightRecord.recorded_date)
        ):
            writer.writerow([weight_row.recorded_date, weight_row.weight_kg, weight_row.note or ""])
    elif kind == "diary":
        writer.writerow(["occurred_at", "title", "category", "severity", "body"])
        for diary_row in db.scalars(
            select(DiaryEntry)
            .where(DiaryEntry.patient_id == patient_id)
            .order_by(DiaryEntry.occurred_at)
        ):
            writer.writerow(
                [diary_row.occurred_at, diary_row.title or "", diary_row.category, diary_row.severity or "", diary_row.body]
            )
    else:
        raise HTTPException(status_code=404, detail="Unknown export")
    payload = io.BytesIO(output.getvalue().encode("utf-8"))
    return StreamingResponse(
        payload,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{kind}.csv"'},
    )
